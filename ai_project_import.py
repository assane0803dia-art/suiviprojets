"""
Import de projet depuis un document Word/PDF, via extraction structurée par IA.

Principe : le document décrivant un projet (souvent un document de projet, une
proposition, ou un rapport) est lu, son texte extrait, puis envoyé à l'IA avec
des instructions strictes pour qu'elle retourne une structure JSON fidèle au
contenu — jamais des valeurs inventées pour "compléter" ce qui manque.

L'écran de vérification (dans la page d'import) reste l'étape non-négociable
avant tout enregistrement en base : cette extraction ne doit jamais écrire
directement dans la base de données sans passage par une relecture humaine.
"""
import io
import json
import re

import streamlit as st
import anthropic
from docx import Document as DocxDocument
from pypdf import PdfReader


# ------------------------------------------------------------------------
# Extraction du texte brut depuis le fichier déposé
# ------------------------------------------------------------------------
def extraire_texte_document(uploaded_file) -> str:
    """Extrait le texte brut d'un fichier .docx ou .pdf déposé via st.file_uploader."""
    nom = uploaded_file.name.lower()
    contenu = uploaded_file.read()

    if nom.endswith(".docx"):
        doc = DocxDocument(io.BytesIO(contenu))
        morceaux = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                ligne = " | ".join(cell.text.strip() for cell in row.cells)
                if ligne.strip(" |"):
                    morceaux.append(ligne)
        return "\n".join(morceaux)

    elif nom.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(contenu))
        morceaux = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(morceaux)

    else:
        raise ValueError(f"Format de fichier non pris en charge : {uploaded_file.name}. Utilisez un fichier .docx ou .pdf.")


# ------------------------------------------------------------------------
# Extraction de la structure via l'API Anthropic
# ------------------------------------------------------------------------
SCHEMA_ATTENDU = """{
  "projet": {
    "nom": "string",
    "description": "string ou null",
    "date_debut": "YYYY-MM-DD ou null",
    "date_fin": "YYYY-MM-DD ou null",
    "budget": nombre ou null,
    "devise": "FCFA, EUR, USD... ou null"
  },
  "objectifs": [
    {
      "type_objectif": "Général ou Spécifique",
      "titre": "string",
      "description": "string ou null",
      "resultats": [
        {
          "titre": "string",
          "description": "string ou null",
          "indicateur": "string ou null",
          "baseline": nombre ou null,
          "valeur_cible": nombre ou null,
          "unite": "string ou null",
          "activites": [
            {
              "titre": "string",
              "description": "string ou null",
              "date_debut": "YYYY-MM-DD ou null",
              "date_fin": "YYYY-MM-DD ou null",
              "budget": nombre ou null
            }
          ]
        }
      ]
    }
  ]
}"""


def _construire_prompt(texte_document: str) -> str:
    return f"""Tu es un assistant d'extraction de données pour un outil de gestion de projet.
On te fournit le texte d'un document décrivant un projet (document de projet, proposition,
rapport de suivi...). Ta tâche : extraire sa structure en JSON, selon le schéma ci-dessous.

RÈGLES STRICTES (le plus important) :
- N'INVENTE JAMAIS une valeur qui n'est pas explicitement présente dans le document.
- Si une information est absente ou ambiguë, mets `null` — ne devine pas, n'estime pas,
  ne complète pas "pour que ça ait l'air complet".
- Les nombres (budget, baseline, cible) doivent être des valeurs EXACTES du document,
  jamais des arrondis ou des approximations de ta part.
- Si le document ne présente pas clairement d'objectifs/résultats/activités séparés,
  fais de ton mieux pour structurer selon ce qui est écrit, sans inventer de hiérarchie
  qui ne serait pas dans le texte.
- Réponds UNIQUEMENT avec le JSON, sans texte avant ou après, sans balises markdown.

SCHÉMA ATTENDU :
{SCHEMA_ATTENDU}

TEXTE DU DOCUMENT :
---
{texte_document[:100000]}
---

Réponds uniquement avec le JSON."""


def extraire_structure_projet(texte_document: str, model: str = "claude-sonnet-5") -> dict:
    """
    Envoie le texte du document à l'IA et retourne la structure extraite (dict).
    Lève une RuntimeError si la clé API manque, ou une ValueError si la réponse
    n'est pas un JSON valide (mieux vaut échouer clairement que de deviner).
    """
    api_key = st.secrets.get("ANTHROPIC_API_KEY")
    if not api_key:
        raise RuntimeError("Clé API Anthropic manquante. Ajoutez ANTHROPIC_API_KEY dans .streamlit/secrets.toml.")

    if not texte_document or not texte_document.strip():
        raise ValueError("Le document est vide ou son texte n'a pas pu être extrait (fichier scanné/image sans texte ?).")

    client = anthropic.Anthropic(api_key=api_key)
    response = client.messages.create(
        model=model,
        max_tokens=8000,
        messages=[{"role": "user", "content": _construire_prompt(texte_document)}],
    )
    texte_reponse = "".join(block.text for block in response.content if hasattr(block, "text"))

    # Au cas où l'IA ajouterait malgré tout des balises markdown autour du JSON
    texte_nettoye = re.sub(r"^```(?:json)?\s*|\s*```$", "", texte_reponse.strip())

    try:
        return json.loads(texte_nettoye)
    except json.JSONDecodeError:
        # Erreur la plus fréquente : une virgule superflue juste avant une accolade/crochet
        # fermant (ex: dernier élément d'une liste) — on tente une réparation ciblée avant
        # d'abandonner, plutôt que d'échouer sur un défaut aussi mineur et fréquent.
        texte_repare = re.sub(r",(\s*[}\]])", r"\1", texte_nettoye)
        try:
            return json.loads(texte_repare)
        except json.JSONDecodeError as e:
            # Conserve la réponse brute pour que l'écran d'import puisse l'afficher —
            # utile pour comprendre ce qui a précisément posé problème.
            st.session_state["import_derniere_reponse_brute"] = texte_reponse
            raise ValueError(
                f"L'IA n'a pas retourné un JSON valide — réessayez, ou le document est peut-être "
                f"trop désorganisé pour une extraction automatique. Détail technique : {e}"
            )